# -*- coding: utf-8 -*-
################################################################################
#
#    AllyRoot Tech
#
#    Copyright (C) 2026-TODAY AllyRoot Tech (allyroottech@gmail.com).
#
#    You can modify it under the terms of the GNU LESSER
#    GENERAL PUBLIC LICENSE (LGPL v3), Version 3.
#
#    This program is distributed in the hope that it will be useful,
#    but WITHOUT ANY WARRANTY; without even the implied warranty of
#    MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
#    GNU LESSER GENERAL PUBLIC LICENSE (LGPL v3) for more details.
#
#    You should have received a copy of the GNU LESSER GENERAL PUBLIC LICENSE
#    (LGPL v3) along with this program.
#    If not, see <http://www.gnu.org/licenses/>.
#
################################################################################

import csv
import importlib.util
import io
import json
import logging
import mimetypes
import os
import shutil
import subprocess
import tarfile
import tempfile
import time
import zipfile
from urllib.parse import quote

from markupsafe import Markup

from odoo import http
from odoo.addons.mail.tools.discuss import add_guest_to_context
from odoo.http import content_disposition, request

_logger = logging.getLogger(__name__)

_PREVIEW_TEMPLATE = "chatter_attachment_preview.preview_page"

# Excel preview and Office-document conversion rely on optional software most
# Odoo installs won't have out of the box. None of it is required for the
# module to install, or for the rest of the preview types (images, video,
# audio, PDF, CSV, ZIP/TAR, markdown, text) to work — so it is deliberately
# *not* declared under `external_dependencies` in the manifest, which would
# block installation entirely for anyone missing it. Instead we only warn at
# load time, the same way core Odoo does for other optional libraries (see
# addons/attachment_indexation/models/ir_attachment.py and its pdfminer.six
# check). This tells the admin what to install right when the module is
# installed/upgraded, without ever surfacing a traceback to an end user.
_HAS_OPENPYXL = importlib.util.find_spec("openpyxl") is not None
_HAS_XLRD = importlib.util.find_spec("xlrd") is not None
_HAS_LIBREOFFICE = bool(shutil.which("libreoffice") or shutil.which("soffice"))

if not _HAS_OPENPYXL:
    _logger.warning(
        "chatter_attachment_preview: Python package 'openpyxl' is not installed — "
        ".xlsx attachments will show a download-only notice instead of an inline "
        "preview. Run `pip install openpyxl` to enable it."
    )
if not _HAS_XLRD:
    _logger.warning(
        "chatter_attachment_preview: Python package 'xlrd' is not installed — "
        "legacy .xls attachments will show a download-only notice instead of an "
        "inline preview. Run `pip install xlrd` to enable it."
    )
if not _HAS_LIBREOFFICE:
    _logger.warning(
        "chatter_attachment_preview: LibreOffice was not found on PATH (checked "
        "'libreoffice' and 'soffice') — Office documents (.docx, .pptx, .odt, ...) "
        "will show a download-only notice instead of an inline preview. Install "
        "LibreOffice to enable it."
    )


class ChatterAttachmentPreviewController(http.Controller):

    # 2 MB cap for inline text/markdown previews — keeps the iframe responsive
    # and bounds how much attachment content we ever decode and embed as JSON.
    _TEXT_PREVIEW_MAX_BYTES = 2 * 1024 * 1024
    # 20 MB cap on a single decompressed ZIP/TAR entry. Checked both against
    # the archive's declared size *and* while reading, so a crafted archive
    # that under-reports its size can't be used to exhaust server memory
    # (a "zip bomb").
    _ENTRY_PREVIEW_MAX_BYTES = 20 * 1024 * 1024

    def _get_attachment(self, attachment_id, access_token=None):
        try:
            attachment = request.env["ir.binary"]._find_record(
                res_model="ir.attachment",
                res_id=int(attachment_id),
                access_token=access_token,
                field="raw",
            )
        except Exception as exc:
            raise request.not_found() from exc
        return attachment.sudo()

    def _content_url(self, attachment, access_token=None, download=False):
        params = {"filename": attachment.name or ""}
        if access_token:
            params["access_token"] = access_token
        if download:
            params["download"] = "true"
        query = "&".join(
            f"{quote(str(k))}={quote(str(v))}" for k, v in params.items() if v
        )
        return f"/web/content/{attachment.id}?{query}" if query else f"/web/content/{attachment.id}"

    def _image_url(self, attachment, access_token=None):
        params = {}
        if access_token:
            params["access_token"] = access_token
        if attachment.name:
            params["filename"] = attachment.name
        query = "&".join(
            f"{quote(str(k))}={quote(str(v))}" for k, v in params.items() if v
        )
        return f"/web/image/{attachment.id}?{query}" if query else f"/web/image/{attachment.id}"

    @staticmethod
    def _fmt_size(n):
        for unit in ("B", "KB", "MB", "GB"):
            if n < 1024:
                return f"{n} {unit}" if unit == "B" else f"{n:.1f} {unit}"
            n /= 1024
        return f"{n:.1f} TB"

    def _read_capped(self, fileobj, cap):
        """Read at most `cap` bytes from `fileobj`. Returns (data, exceeded) —
        exceeded is True when more than `cap` bytes were available, in which
        case `data` is empty (callers show a notice rather than a truncated,
        possibly-corrupt binary)."""
        chunk = fileobj.read(cap + 1)
        if len(chunk) > cap:
            return b"", True
        return chunk, False

    def _is_raw_zip(self, attachment):
        """True only for plain ZIP archives, not OOXML/ODF documents.
        OOXML (docx/xlsx/pptx) always contain [Content_Types].xml.
        ODF (odt/ods/odp) always contain a 'mimetype' entry."""
        content = attachment.raw or b""
        try:
            buf = io.BytesIO(content)
            if not zipfile.is_zipfile(buf):
                return False
            buf.seek(0)
            with zipfile.ZipFile(buf) as zf:
                names = set(zf.namelist())
                if "[Content_Types].xml" in names:
                    return False
                if "mimetype" in names:
                    return False
        except zipfile.BadZipFile:
            return False
        return True

    def _zip_data(self, attachment, url_prefix, max_entries=1000):
        content = attachment.raw or b""
        entries, truncated = [], False
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as zf:
                members = zf.infolist()
                total = len(members)
                for info in members[:max_entries]:
                    depth = max(0, info.filename.rstrip("/").count("/"))
                    is_dir = info.filename.endswith("/")
                    entries.append({
                        "name": info.filename,
                        "basename": info.filename.rstrip("/").rsplit("/", 1)[-1] or info.filename,
                        "size": self._fmt_size(info.file_size) if not is_dir else "",
                        "is_dir": is_dir,
                        "depth": depth,
                        "date": "%04d-%02d-%02d" % info.date_time[:3] if info.date_time[0] > 1980 else "",
                        "entry_url": "" if is_dir else (url_prefix + quote(info.filename, safe="")),
                    })
                truncated = total > max_entries
        except zipfile.BadZipFile:
            pass
        return {
            "zip_entries": entries,
            "zip_count": len([e for e in entries if not e["is_dir"]]),
            "zip_dir_count": len([e for e in entries if e["is_dir"]]),
            "zip_truncated": truncated,
            "zip_max": max_entries,
        }

    def _is_tar(self, attachment):
        try:
            with tarfile.open(fileobj=io.BytesIO(attachment.raw or b"")):
                pass
            return True
        except tarfile.TarError:
            return False

    def _tar_data(self, attachment, url_prefix, max_entries=1000):
        content = attachment.raw or b""
        entries, truncated = [], False
        try:
            with tarfile.open(fileobj=io.BytesIO(content)) as tf:
                members = tf.getmembers()
                total = len(members)
                for m in members[:max_entries]:
                    is_dir = m.isdir()
                    depth = max(0, m.name.count("/"))
                    basename = m.name.rsplit("/", 1)[-1] or m.name
                    try:
                        date = time.strftime("%Y-%m-%d", time.gmtime(m.mtime)) if m.mtime else ""
                    except (OSError, OverflowError):
                        date = ""
                    entries.append({
                        "name": m.name,
                        "basename": basename,
                        "size": self._fmt_size(m.size) if not is_dir else "",
                        "is_dir": is_dir,
                        "depth": depth,
                        "date": date,
                        "entry_url": "" if is_dir else (url_prefix + quote(m.name, safe="")),
                    })
                truncated = total > max_entries
        except tarfile.TarError:
            pass
        return {
            "zip_entries": entries,
            "zip_count": len([e for e in entries if not e["is_dir"]]),
            "zip_dir_count": len([e for e in entries if e["is_dir"]]),
            "zip_truncated": truncated,
            "zip_max": max_entries,
        }

    def _is_text(self, attachment):
        raw = attachment.raw or b""
        if not raw:
            return False
        sample = raw[:8192]
        try:
            sample.decode("utf-8")
            return b"\x00" not in sample
        except (UnicodeDecodeError, ValueError):
            return False

    def _is_markdown(self, attachment):
        """Detect markdown by mimetype or by filename extension (.md / .markdown)."""
        if "markdown" in (attachment.mimetype or ""):
            return True
        return (attachment.name or "").lower().endswith((".md", ".markdown"))

    def _spreadsheet_format(self, attachment):
        """Detect spreadsheet sub-format from the attachment mimetype. Returns 'csv', 'xls',
        or 'xlsx', or None when the mimetype is not a Python-parseable spreadsheet."""
        mt = attachment.mimetype or ""
        if mt == "text/csv":
            return "csv"
        if "ms-excel" in mt:
            return "xls"
        if "spreadsheetml" in mt:
            return "xlsx"
        return None

    def _column_label(self, index):
        label = ""
        while index:
            index, remainder = divmod(index - 1, 26)
            label = chr(65 + remainder) + label
        return label

    def _spreadsheet_data(self, attachment, fmt, max_rows=300, max_cols=60):
        rows, sheet_name, truncated = [], "", False
        content = attachment.raw or b""

        if fmt == "xlsx":
            import openpyxl  # noqa: PLC0415
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                sheet_name = wb.sheetnames[0] if wb.sheetnames else ""
                ws = wb[sheet_name] if sheet_name else None
                if ws:
                    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
                        rows.append(["" if v is None else str(v) for v in row[:max_cols]])
                        if i >= max_rows:
                            truncated = True
                            break
            finally:
                wb.close()
        elif fmt == "xls":
            import xlrd  # noqa: PLC0415
            wb = xlrd.open_workbook(file_contents=content)
            ws = wb.sheet_by_index(0)
            sheet_name = ws.name
            for i in range(min(ws.nrows, max_rows)):
                rows.append([str(v) for v in ws.row_values(i, 0, min(ws.ncols, max_cols))])
            truncated = ws.nrows > max_rows or ws.ncols > max_cols
        elif fmt == "csv":
            sheet_name = "CSV"
            for i, row in enumerate(
                csv.reader(io.StringIO(content.decode("utf-8-sig", errors="replace"))), 1
            ):
                rows.append([str(v) for v in row[:max_cols]])
                if i >= max_rows:
                    truncated = True
                    break

        col_count = max((len(r) for r in rows), default=0)
        for row in rows:
            if len(row) < col_count:
                row.extend([""] * (col_count - len(row)))

        return {
            "sheet_name": sheet_name or "Sheet 1",
            "rows": rows,
            "col_headers": ["#"] + [self._column_label(i) for i in range(1, col_count + 1)],
            "row_count": len(rows),
            "col_count": col_count,
            "truncated": truncated,
            "max_rows": max_rows,
            "max_cols": max_cols,
        }

    def _render_preview(self, values):
        response = request.render(_PREVIEW_TEMPLATE, values)
        response.headers["X-Frame-Options"] = "SAMEORIGIN"
        return response

    @http.route(
        "/chatter_attachment_preview/preview/<int:attachment_id>",
        type="http", auth="public", methods=["GET"], readonly=True,
    )
    @add_guest_to_context
    def preview(self, attachment_id, access_token=None, **kwargs):
        attachment = self._get_attachment(attachment_id, access_token=access_token)
        mimetype = attachment.mimetype or "application/octet-stream"
        values = {
            "name": attachment.name or "Attachment",
            "content_url": self._content_url(attachment, access_token=access_token),
            "download_url": self._content_url(attachment, access_token=access_token, download=True),
        }

        # `_detect_preview` can raise (corrupt file, missing optional library,
        # …) — it never returns a broken/partial preview. Any failure here is
        # turned into a clean "notice" card with a Download fallback instead
        # of ever bubbling up into Odoo's generic error page inside the iframe.
        try:
            self._detect_preview(attachment, mimetype, access_token, values)
        except ImportError as exc:
            modname = getattr(exc, "name", None) or str(exc)
            _logger.warning(
                "chatter_attachment_preview: missing optional dependency %r while "
                "previewing attachment %s — showing a notice instead.",
                modname, attachment.id,
            )
            values["preview_type"] = "notice"
            values["notice_title"] = "Preview requires an additional library"
            values["notice_body"] = (
                f'The server is missing the Python package "{modname}" needed to '
                "preview this file. Ask your administrator to install it, or "
                "download the file instead."
            )
        except Exception:
            _logger.exception(
                "chatter_attachment_preview: failed to build preview for attachment %s",
                attachment.id,
            )
            values["preview_type"] = "notice"
            values["notice_title"] = "This file cannot be previewed"
            values["notice_body"] = (
                "Something went wrong while generating the preview. "
                "You can download the file instead."
            )

        return self._render_preview(values)

    def _detect_preview(self, attachment, mimetype, access_token, values):
        """Detect the attachment's preview type and populate `values` with
        whatever the matching template needs. Raises on failure — see
        `preview()`, which is the only caller and always handles it."""
        if mimetype.startswith("image/"):
            values["preview_type"] = "image"
            values["image_url"] = self._image_url(attachment, access_token=access_token)
        elif mimetype.startswith("video/"):
            values["preview_type"] = "video"
        elif mimetype.startswith("audio/"):
            values["preview_type"] = "audio"
        elif mimetype == "application/pdf":
            values["preview_type"] = "pdf"
            values["pdf_viewer_url"] = (
                "/web/static/lib/pdfjs/web/viewer.html?file="
                + quote(request.httprequest.host_url.rstrip("/") + values["content_url"], safe="")
                + "#pagemode=none"
            )
        else:
            fmt = self._spreadsheet_format(attachment)
            if fmt:
                values["preview_type"] = "spreadsheet"
                values.update(self._spreadsheet_data(attachment, fmt))
            elif self._is_markdown(attachment):
                raw = attachment.raw or b""
                text = raw[:self._TEXT_PREVIEW_MAX_BYTES].decode("utf-8", errors="replace")
                if len(raw) > self._TEXT_PREVIEW_MAX_BYTES:
                    text += "\n\n… (preview truncated at 2 MB)"
                # Escape <, >, & so the JSON is safe to embed inside a <script> tag.
                safe_json = (
                    json.dumps(text)
                    .replace("<", "\\u003c")
                    .replace(">", "\\u003e")
                    .replace("&", "\\u0026")
                )
                values["preview_type"] = "markdown"
                values["markdown_json"] = Markup(safe_json)
            elif self._is_raw_zip(attachment):
                base = f"/chatter_attachment_preview/zip_entry/{attachment.id}"
                url_prefix = (
                    f"{base}?access_token={quote(access_token)}&path="
                    if access_token else f"{base}?path="
                )
                values["preview_type"] = "zip"
                values["zip_back_url"] = values["content_url"]
                values.update(self._zip_data(attachment, url_prefix))
            elif self._is_tar(attachment):
                base = f"/chatter_attachment_preview/tar_entry/{attachment.id}"
                url_prefix = (
                    f"{base}?access_token={quote(access_token)}&path="
                    if access_token else f"{base}?path="
                )
                values["preview_type"] = "zip"
                values["zip_back_url"] = values["content_url"]
                values.update(self._tar_data(attachment, url_prefix))
            elif self._is_text(attachment):
                raw = attachment.raw or b""
                text = raw[:self._TEXT_PREVIEW_MAX_BYTES].decode("utf-8", errors="replace")
                if len(raw) > self._TEXT_PREVIEW_MAX_BYTES:
                    text += "\n\n… (preview truncated at 2 MB)"
                ext = os.path.splitext(attachment.name or "")[1].lstrip(".").lower()
                safe_json = (
                    json.dumps(text)
                    .replace("<", "\\u003c")
                    .replace(">", "\\u003e")
                    .replace("&", "\\u0026")
                )
                values["preview_type"] = "text"
                values["text_json"] = Markup(safe_json)
                values["text_lang"] = ext
                values["back_url"] = ""
            else:
                # Office documents (docx, pptx, odt, rtf, …) — LibreOffice converts to PDF.
                converted_url = f"/chatter_attachment_preview/pdf/{attachment.id}"
                if access_token:
                    converted_url += f"?access_token={quote(access_token)}"
                values["preview_type"] = "office"
                values["converted_url"] = converted_url

    @http.route(
        "/chatter_attachment_preview/zip_entry/<int:attachment_id>",
        type="http", auth="public", methods=["GET"], readonly=True,
    )
    @add_guest_to_context
    def zip_entry(self, attachment_id, path="", access_token=None, **kwargs):
        attachment = self._get_attachment(attachment_id, access_token=access_token)
        back_url = f"/chatter_attachment_preview/preview/{attachment_id}"
        if access_token:
            back_url += f"?access_token={quote(access_token)}"

        try:
            with zipfile.ZipFile(io.BytesIO(attachment.raw or b"")) as zf:
                info = zf.getinfo(path)
                if info.file_size > self._ENTRY_PREVIEW_MAX_BYTES:
                    entry_bytes, too_large = b"", True
                else:
                    with zf.open(info) as fh:
                        entry_bytes, too_large = self._read_capped(fh, self._ENTRY_PREVIEW_MAX_BYTES)
        except (zipfile.BadZipFile, KeyError):
            raise request.not_found()

        name = path.rsplit("/", 1)[-1] or path

        if too_large:
            return self._render_preview({
                "name": name,
                "preview_type": "notice",
                "notice_title": "File too large to preview",
                "notice_body": (
                    f"This file is larger than {self._fmt_size(self._ENTRY_PREVIEW_MAX_BYTES)} "
                    "— open the archive to download it."
                ),
                "content_url": "",
                "download_url": back_url,
            })

        guessed, _ = mimetypes.guess_type(name)
        mimetype = guessed or "application/octet-stream"

        if mimetype.startswith("image/"):
            return request.make_response(entry_bytes, [
                ("Content-Type", mimetype),
                ("Content-Length", str(len(entry_bytes))),
            ])

        # Try to read as text (explicit text types + heuristic UTF-8 check)
        text = None
        if (
            mimetype.startswith("text/")
            or mimetype in ("application/json", "application/javascript",
                            "application/xml", "application/x-sh")
        ):
            text = entry_bytes.decode("utf-8", errors="replace")
        else:
            try:
                decoded = entry_bytes.decode("utf-8")
                if "\x00" not in decoded:
                    text = decoded
            except (UnicodeDecodeError, ValueError):
                pass

        if text is not None:
            if len(text) > self._TEXT_PREVIEW_MAX_BYTES:
                text = text[:self._TEXT_PREVIEW_MAX_BYTES] + "\n\n… (preview truncated at 2 MB)"
            ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
            safe_json = (
                json.dumps(text)
                .replace("<", "\\u003c")
                .replace(">", "\\u003e")
                .replace("&", "\\u0026")
            )
            return self._render_preview({
                "name": name,
                "preview_type": "text",
                "text_json": Markup(safe_json),
                "text_lang": ext,
                "back_url": back_url,
                "content_url": "",
                "download_url": "",
            })

        return self._render_preview({
            "name": name,
            "preview_type": "notice",
            "notice_title": "Cannot preview this file",
            "notice_body": f"Binary file ({mimetype}) — open the ZIP to download it.",
            "content_url": "",
            "download_url": back_url,
        })

    @http.route(
        "/chatter_attachment_preview/tar_entry/<int:attachment_id>",
        type="http", auth="public", methods=["GET"], readonly=True,
    )
    @add_guest_to_context
    def tar_entry(self, attachment_id, path="", access_token=None, **kwargs):
        attachment = self._get_attachment(attachment_id, access_token=access_token)
        back_url = f"/chatter_attachment_preview/preview/{attachment_id}"
        if access_token:
            back_url += f"?access_token={quote(access_token)}"

        try:
            with tarfile.open(fileobj=io.BytesIO(attachment.raw or b"")) as tf:
                member = tf.getmember(path)
                if member.size > self._ENTRY_PREVIEW_MAX_BYTES:
                    entry_bytes, too_large = b"", True
                else:
                    fobj = tf.extractfile(member)
                    if fobj is None:
                        entry_bytes, too_large = b"", False
                    else:
                        with fobj:
                            entry_bytes, too_large = self._read_capped(fobj, self._ENTRY_PREVIEW_MAX_BYTES)
        except (tarfile.TarError, KeyError):
            raise request.not_found()

        name = path.rsplit("/", 1)[-1] or path

        if too_large:
            return self._render_preview({
                "name": name,
                "preview_type": "notice",
                "notice_title": "File too large to preview",
                "notice_body": (
                    f"This file is larger than {self._fmt_size(self._ENTRY_PREVIEW_MAX_BYTES)} "
                    "— download the archive to access it."
                ),
                "content_url": "",
                "download_url": back_url,
            })

        guessed, _ = mimetypes.guess_type(name)
        mimetype = guessed or "application/octet-stream"

        if mimetype.startswith("image/"):
            return request.make_response(entry_bytes, [
                ("Content-Type", mimetype),
                ("Content-Length", str(len(entry_bytes))),
            ])

        text = None
        if mimetype.startswith("text/") or mimetype in (
            "application/json", "application/javascript",
            "application/xml", "application/x-sh",
        ):
            text = entry_bytes.decode("utf-8", errors="replace")
        else:
            try:
                decoded = entry_bytes.decode("utf-8")
                if "\x00" not in decoded:
                    text = decoded
            except (UnicodeDecodeError, ValueError):
                pass

        if text is not None:
            if len(text) > self._TEXT_PREVIEW_MAX_BYTES:
                text = text[:self._TEXT_PREVIEW_MAX_BYTES] + "\n\n… (preview truncated at 2 MB)"
            ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
            safe_json = (
                json.dumps(text)
                .replace("<", "\\u003c")
                .replace(">", "\\u003e")
                .replace("&", "\\u0026")
            )
            return self._render_preview({
                "name": name,
                "preview_type": "text",
                "text_json": Markup(safe_json),
                "text_lang": ext,
                "back_url": back_url,
                "content_url": "",
                "download_url": "",
            })

        return self._render_preview({
            "name": name,
            "preview_type": "notice",
            "notice_title": "Cannot preview this file",
            "notice_body": f"Binary file ({mimetype}) — download the archive to access it.",
            "content_url": "",
            "download_url": back_url,
        })

    @http.route(
        "/chatter_attachment_preview/pdf/<int:attachment_id>",
        type="http", auth="public", methods=["GET"], readonly=True,
    )
    @add_guest_to_context
    def converted_pdf(self, attachment_id, access_token=None, **kwargs):
        attachment = self._get_attachment(attachment_id, access_token=access_token)
        soffice = shutil.which("libreoffice") or shutil.which("soffice")
        if not soffice:
            return self._render_preview({
                "name": attachment.name or "Attachment",
                "preview_type": "notice",
                "notice_title": "Document converter is not installed",
                "notice_body": "Install LibreOffice on the Odoo server to preview Office documents.",
                "content_url": self._content_url(attachment, access_token=access_token),
                "download_url": self._content_url(attachment, access_token=access_token, download=True),
            })

        # Use the file's own extension so LibreOffice identifies the source format correctly.
        ext = os.path.splitext(attachment.name or "")[1].lstrip(".") or "bin"
        with tempfile.TemporaryDirectory(prefix="odoo_cap_") as tmpdir:
            src = os.path.join(tmpdir, f"source.{ext}")
            try:
                with open(src, "wb") as f:
                    f.write(attachment.raw or b"")
                subprocess.run(
                    [soffice, "--headless", "--convert-to", "pdf", "--outdir", tmpdir, src],
                    check=True, timeout=45,
                    stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                )
            except (subprocess.SubprocessError, OSError):
                return self._render_preview({
                    "name": attachment.name or "Attachment",
                    "preview_type": "notice",
                    "notice_title": "This file type cannot be previewed",
                    "notice_body": "This file cannot be converted for preview. You can download it instead.",
                    "content_url": self._content_url(attachment, access_token=access_token),
                    "download_url": self._content_url(attachment, access_token=access_token, download=True),
                })
            pdf_path = os.path.join(tmpdir, "source.pdf")
            if not os.path.exists(pdf_path):
                return self._render_preview({
                    "name": attachment.name or "Attachment",
                    "preview_type": "notice",
                    "notice_title": "This file type cannot be previewed",
                    "notice_body": "This file cannot be converted for preview. You can download it instead.",
                    "content_url": self._content_url(attachment, access_token=access_token),
                    "download_url": self._content_url(attachment, access_token=access_token, download=True),
                })
            with open(pdf_path, "rb") as f:
                pdf_content = f.read()

        return request.make_response(pdf_content, [
            ("Content-Type", "application/pdf"),
            ("Content-Length", str(len(pdf_content))),
            ("Content-Disposition", content_disposition(f"{attachment.name or 'attachment'}.pdf", "inline")),
        ])
